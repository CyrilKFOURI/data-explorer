"""Turn one or more source files into a single DuckDB query."""

from pathlib import Path

import duckdb
import openpyxl
import pyarrow as pa

from cyfouri_data_explorer.ingestion.csv_sniff import sniff_csv
from cyfouri_data_explorer.ingestion.options import CsvOptions, SourceOptions
from cyfouri_data_explorer.ingestion.sources import SourceFile
from cyfouri_data_explorer.ingestion.sql import ident, lit

FULL_SNIFF_MAX_BYTES = 100 * 1024 * 1024


def _column(values: list[object]) -> pa.Array:
    """Typed when the cells agree (numbers, dates, booleans), text when they are mixed."""
    try:
        return pa.array(values)
    except (pa.ArrowInvalid, pa.ArrowTypeError):
        return pa.array([None if value is None else str(value) for value in values], pa.string())


def read_sheet(path: Path, sheet: str) -> pa.Table:
    """A sheet as a table: the first row names the columns; fully empty rows are skipped."""
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        rows = [
            row
            for row in workbook[sheet].iter_rows(values_only=True)
            if any(c is not None for c in row)
        ]
    finally:
        workbook.close()
    if not rows:
        return pa.table({})
    width = max(len(row) for row in rows)
    names: list[str] = []
    for index in range(width):
        header = rows[0][index] if index < len(rows[0]) else None
        name = (
            str(header).strip()
            if header is not None and str(header).strip()
            else f"column_{index + 1}"
        )
        while name in names:  # duplicated header: keep both columns
            name = f"{name}_{index + 1}"
        names.append(name)
    body = [list(row) + [None] * (width - len(row)) for row in rows[1:]]
    return pa.table({name: _column([row[i] for row in body]) for i, name in enumerate(names)})


def excel_sheets(path: Path) -> list[str]:
    workbook = openpyxl.load_workbook(path, read_only=True)
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


class SourceReader:
    """Registers what it needs in a DuckDB connection and returns SQL selecting the data."""

    def __init__(self, connection: duckdb.DuckDBPyConnection, options: SourceOptions) -> None:
        self.connection = connection
        self.options = options
        self.excel_rows: dict[Path, int] = {}
        self.csv_options: dict[Path, CsvOptions] = {}
        self._views = 0

    def csv_options_for(self, file: SourceFile) -> CsvOptions:
        if file.path not in self.csv_options:
            self.csv_options[file.path] = self.options.csv or sniff_csv(file.path)
        return self.csv_options[file.path]

    def csv_sql(self, file: SourceFile) -> str:
        options = self.csv_options_for(file)
        sample_size = -1 if file.size <= FULL_SNIFF_MAX_BYTES else 200_000
        return (
            f"SELECT * FROM read_csv({lit(str(file.path))}, delim={lit(options.delimiter)}, "
            f"header={str(options.header).lower()}, decimal_separator={lit(options.decimal)}, "
            f"encoding={lit(options.encoding)}, sample_size={sample_size})"
        )

    def excel_sql(self, file: SourceFile, sheets: list[str] | None) -> str:
        names = sheets or excel_sheets(file.path)[:1]
        parts = []
        self.excel_rows[file.path] = 0
        for sheet in names:
            table = read_sheet(file.path, sheet)
            self._views += 1
            view = f"excel_sheet_{self._views}"
            self.connection.register(view, table)
            self.excel_rows[file.path] += table.num_rows
            sheet_column = f", {lit(sheet)} AS source_sheet" if len(names) > 1 else ""
            parts.append(f"SELECT *{sheet_column} FROM {view}")
        return " UNION ALL BY NAME ".join(parts)

    def file_sql(self, file: SourceFile, sheets: list[str] | None = None) -> str:
        if file.format == "csv":
            return self.csv_sql(file)
        if file.format == "parquet":
            return f"SELECT * FROM read_parquet({lit(str(file.path))})"
        return self.excel_sql(file, sheets)

    def columns(self, sql: str) -> list[tuple[str, str]]:
        """(name, DuckDB type) of each column of a query."""
        rows = self.connection.execute(f"DESCRIBE ({sql})").fetchall()
        return [(row[0], row[1]) for row in rows]

    def combined_sql(self, files: list[SourceFile]) -> str:
        if len(files) == 1:
            return self.file_sql(files[0], self.options.sheets)
        parts = [
            f"SELECT *, {lit(file.path.name)} AS source_file FROM ({self.file_sql(file)})"
            for file in files
        ]
        if self.options.folder_mode == "union":
            return " UNION ALL BY NAME ".join(parts)
        part_columns = [[name for name, _ in self.columns(part)] for part in parts]
        common = [name for name in part_columns[0] if all(name in cols for cols in part_columns)]
        selected = ", ".join(ident(name) for name in common)
        return " UNION ALL ".join(f"SELECT {selected} FROM ({part})" for part in parts)

    def schema_mismatch(self, files: list[SourceFile]) -> bool:
        if len(files) < 2:
            return False
        names = [{name for name, _ in self.columns(self.file_sql(file))} for file in files]
        return any(other != names[0] for other in names[1:])
