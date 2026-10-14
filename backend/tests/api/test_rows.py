import io
import json
import time
from pathlib import Path

import duckdb
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.config import default_samples_dir
from app.domain.rows import export as export_module
from tests.api.test_profiling import import_path


@pytest.fixture
def traps(client: TestClient) -> str:
    return import_path(client, default_samples_dir() / "typing_traps.csv")


@pytest.fixture
def twins(client: TestClient, tmp_path: Path) -> str:
    path = tmp_path / "twins.csv"
    path.write_text("name,amount\nana,1\nbob,2\nana,1\ncid,\nbob,3\n", encoding="utf-8")
    return import_path(client, path)


def url(dataset_id: str, path: str = "rows") -> str:
    return f"/api/v1/datasets/{dataset_id}/{path}"


def test_rows_of_a_dataset(client: TestClient, traps: str) -> None:
    body = client.get(url(traps), params={"offset": 10, "limit": 5}).json()

    assert body["total"] == 600 and body["offset"] == 10
    assert len(body["rows"]) == 5 and body["rows"][0][0] == 11
    names = [c["name"] for c in body["columns"]]
    assert "customer_id" in names and len(body["rows"][0]) == len(names) + 1


def test_modes_search_sort_and_filters(client: TestClient, twins: str) -> None:
    duplicates = client.get(url(twins), params={"mode": "duplicates"}).json()
    missing = client.get(url(twins), params={"mode": "missing"}).json()
    found = client.get(url(twins), params={"search": "BO", "sort": "amount", "desc": True}).json()
    filtered = client.get(
        url(twins),
        params={"filters": json.dumps([{"column": "name", "op": "in", "values": ["ana"]}])},
    ).json()

    assert [row[0] for row in duplicates["rows"]] == [1, 3]
    assert [row[0] for row in missing["rows"]] == [4]
    assert [row[0] for row in found["rows"]] == [5, 2]
    assert filtered["total"] == 2


def test_bad_requests(client: TestClient, twins: str) -> None:
    assert client.get(url(twins), params={"sort": "nope"}).status_code == 422
    assert client.get(url(twins), params={"mode": "all"}).status_code == 422
    assert client.get(url(twins), params={"limit": 501}).status_code == 422


def test_exports(client: TestClient, twins: str, tmp_path: Path) -> None:
    as_csv = client.get(url(twins, "export"), params={"format": "csv", "columns": "amount,name"})
    as_parquet = client.get(
        url(twins, "export"), params={"format": "parquet", "mode": "duplicates"}
    )
    as_excel = client.get(url(twins, "export"), params={"format": "xlsx", "search": "cid"})
    summary = client.get(url(twins, "export/summary"))

    assert as_csv.status_code == 200, as_csv.text
    assert "attachment" in as_csv.headers["content-disposition"]
    assert ".csv" in as_csv.headers["content-disposition"]
    assert as_csv.text.splitlines()[:2] == ["amount,name", "1,ana"]
    saved = tmp_path / "out.parquet"
    saved.write_bytes(as_parquet.content)
    assert len(duckdb.sql(f"SELECT * FROM '{saved.as_posix()}'").fetchall()) == 2
    sheet = load_workbook(io.BytesIO(as_excel.content)).active
    assert sheet is not None and sheet.max_row == 2
    assert summary.text.splitlines()[0].startswith("column,type,rows,missing")
    assert len(summary.text.splitlines()) == 3


def test_excel_limit(client: TestClient, twins: str, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(export_module, "EXCEL_ROWS", 3)

    response = client.get(url(twins, "export"), params={"format": "xlsx"})

    assert response.status_code == 422 and response.json()["error"]["code"] == "too_many_rows"


@pytest.mark.slow
def test_five_million_rows(client: TestClient, tmp_path: Path) -> None:
    path = tmp_path / "big.parquet"
    duckdb.execute(
        "COPY (SELECT range AS id, random() AS a, ['x', 'y', 'z'][1 + range % 3] AS c "
        f"FROM range(5000000)) TO '{path.as_posix()}' (FORMAT parquet)"
    )
    dataset_id = import_path(client, path)

    start = time.perf_counter()
    first = client.get(url(dataset_id), params={"limit": 0})
    counted = time.perf_counter() - start
    start = time.perf_counter()
    top = client.get(url(dataset_id), params={"limit": 100})
    opened = time.perf_counter() - start
    start = time.perf_counter()
    deep = client.get(url(dataset_id), params={"offset": 4_000_000, "limit": 100})
    paged = time.perf_counter() - start

    assert first.json()["total"] == 5_000_000
    assert top.json()["rows"][0][0] == 1 and deep.json()["rows"][0][0] == 4_000_001
    assert counted < 1, counted
    assert opened < 0.5, opened
    assert paged < 0.5, paged
